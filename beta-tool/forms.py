"""Outil « Questionnaires » : créer un formulaire, le partager par un lien public
(/f/<id>), recevoir les réponses (JSON Lines dans DATA_DIR/forms) et les exporter."""

import csv
import io
import re
import secrets
import time

from flask import Blueprint, current_app, g, jsonify, request, send_file

import activity
import store

bp = Blueprint("forms", __name__)
TYPES = {"text", "paragraph", "email", "number", "date", "choice", "checkbox", "dropdown", "scale", "rating", "yesno"}
CHOICES = {"choice", "checkbox", "dropdown"}


def forms():
    return store.load("forms", {})


def owned(form_id):
    f = forms().get(form_id)
    if not f or (f["owner"] != g.user["name"] and g.user["role"] != "admin"):
        raise ValueError("Questionnaire introuvable")
    return f


def responses(form_id):
    return store.read_lines(f"forms/{form_id}")


def is_open(f):
    s = f.get("settings", {})
    return s.get("open", True) and not (s.get("closes_at") and time.time() > s["closes_at"])


def clean_question(q):
    qtype = q.get("type") if q.get("type") in TYPES else "text"
    out = {"id": str(q.get("id") or secrets.token_hex(4))[:16], "type": qtype, "label": str(q.get("label") or "")[:300],
           "help": str(q.get("help") or "")[:500], "required": bool(q.get("required"))}
    if qtype in CHOICES:
        out["options"] = [str(o)[:200] for o in (q.get("options") or []) if str(o).strip()][:50] or ["Option 1"]
        out["other"] = bool(q.get("other")) and qtype != "dropdown"
    if qtype == "scale":
        out["min"], out["max"] = int(q.get("min", 1)), max(int(q.get("min", 1)) + 1, min(10, int(q.get("max", 5))))
        out["min_label"], out["max_label"] = str(q.get("min_label") or "")[:60], str(q.get("max_label") or "")[:60]
    return out


def summary(f):
    return {"id": f["id"], "title": f["title"], "created": f["created"], "updated": f.get("updated"), "open": is_open(f),
            "questions": len(f["questions"]), "responses": len(responses(f["id"])), "owner": f["owner"]}


TEMPLATES = {
    "vide": ("Questionnaire sans titre", "", []),
    "sondage": ("Sondage", "Donne ton avis en quelques secondes.", [
        {"type": "choice", "label": "Quelle date te convient le mieux ?", "options": ["Samedi 14", "Dimanche 15", "Samedi 21"], "required": True},
        {"type": "paragraph", "label": "Un commentaire ?"}]),
    "inscription": ("Inscription à l'événement", "Merci de remplir ce formulaire pour t'inscrire.", [
        {"type": "text", "label": "Nom et prénom", "required": True}, {"type": "email", "label": "Adresse e-mail", "required": True},
        {"type": "number", "label": "Nombre de personnes", "required": True}, {"type": "checkbox", "label": "Repas", "options": ["Végétarien", "Sans gluten", "Aucune contrainte"]},
        {"type": "paragraph", "label": "Remarques"}]),
    "avis": ("Ton avis compte", "Aide-nous à nous améliorer.", [
        {"type": "rating", "label": "Note globale", "required": True},
        {"type": "scale", "label": "Recommanderais-tu à un ami ?", "min": 0, "max": 10, "min_label": "Pas du tout", "max_label": "Absolument"},
        {"type": "paragraph", "label": "Qu'est-ce qu'on pourrait améliorer ?"}]),
    "contact": ("Nous contacter", "", [
        {"type": "text", "label": "Nom", "required": True}, {"type": "email", "label": "E-mail", "required": True},
        {"type": "dropdown", "label": "Sujet", "options": ["Question", "Problème", "Autre"]}, {"type": "paragraph", "label": "Message", "required": True}]),
}


# --- Gestion (connecté) ------------------------------------------------------

@bp.get("/api/forms")
def list_forms():
    out = [summary(f) for f in forms().values() if f["owner"] == g.user["name"] or g.user["role"] == "admin"]
    return jsonify(sorted(out, key=lambda f: -(f["updated"] or f["created"])))


@bp.post("/api/forms")
def create():
    tpl = TEMPLATES.get((request.get_json(silent=True) or {}).get("template"), TEMPLATES["vide"])
    fid = secrets.token_urlsafe(6)
    f = {"id": fid, "owner": g.user["name"], "title": tpl[0], "description": tpl[1], "questions": [clean_question(q) for q in tpl[2]],
         "settings": {"open": True, "confirmation": "Merci, ta réponse a bien été enregistrée !", "notify": True, "color": "#1f6bff"},
         "created": time.time(), "updated": time.time()}
    store.update("forms", {}, lambda all_: all_.__setitem__(fid, f))
    return jsonify(f)


@bp.get("/api/forms/<fid>")
def get_form(fid):
    return jsonify({**owned(fid), "responses": len(responses(fid))})


@bp.put("/api/forms/<fid>")
def save(fid):
    owned(fid)
    d = request.get_json(silent=True) or {}

    def go(all_):
        f = all_[fid]
        f["title"] = str(d.get("title") or "Sans titre")[:200]
        f["description"] = str(d.get("description") or "")[:2000]
        f["questions"] = [clean_question(q) for q in (d.get("questions") or [])][:100]
        s = d.get("settings") or {}
        f["settings"] = {"open": bool(s.get("open", True)), "closes_at": float(s["closes_at"]) if s.get("closes_at") else None,
                         "confirmation": str(s.get("confirmation") or "")[:500], "notify": bool(s.get("notify", True)),
                         "color": s.get("color") if re.fullmatch(r"#[0-9a-fA-F]{6}", str(s.get("color"))) else "#1f6bff"}
        f["updated"] = time.time()
        return f
    return jsonify(store.update("forms", {}, go))


@bp.post("/api/forms/<fid>/duplicate")
def duplicate(fid):
    src = owned(fid)
    new_id = secrets.token_urlsafe(6)
    copy = {**src, "id": new_id, "owner": g.user["name"], "title": src["title"] + " (copie)", "created": time.time(), "updated": time.time()}
    store.update("forms", {}, lambda all_: all_.__setitem__(new_id, copy))
    return jsonify(summary(copy))


@bp.delete("/api/forms/<fid>")
def delete(fid):
    owned(fid)
    store.update("forms", {}, lambda all_: all_.pop(fid, None))
    import os
    try:
        os.remove(store.path("forms", f"{fid}.jsonl"))
    except OSError:
        pass
    return jsonify(ok=True)


@bp.get("/api/forms/<fid>/responses")
def list_responses(fid):
    owned(fid)
    return jsonify(responses(fid)[::-1])


@bp.delete("/api/forms/<fid>/responses/<rid>")
def delete_response(fid, rid):
    owned(fid)
    store.rewrite_lines(f"forms/{fid}", [r for r in responses(fid) if r["id"] != rid])
    return jsonify(ok=True)


def flat(q, value):
    if isinstance(value, list):
        return ", ".join(map(str, value))
    return "" if value is None else str(value)


@bp.get("/api/forms/<fid>/export.<fmt>")
def export(fid, fmt):
    f = owned(fid)
    rows = [["Date"] + [q["label"] for q in f["questions"]]]
    for r in responses(fid):
        rows.append([time.strftime("%d/%m/%Y %H:%M", time.localtime(r["at"]))] + [flat(q, r["answers"].get(q["id"])) for q in f["questions"]])
    name = re.sub(r"[^\w\- ]", "", f["title"]).strip() or "reponses"
    if fmt == "csv":
        buf = io.StringIO()
        csv.writer(buf, delimiter=";").writerows(rows)
        data = io.BytesIO(("﻿" + buf.getvalue()).encode("utf-8"))  # BOM : Excel ouvre les accents correctement
        return send_file(data, mimetype="text/csv", as_attachment=True, download_name=f"{name}.csv")
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Réponses"
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F6BFF")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(50, max(12, *(len(str(c.value or "")) for c in col)))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True,
                         download_name=f"{name}.xlsx")
    raise ValueError("Format inconnu")


# --- Pages publiques ----------------------------------------------------------

@bp.get("/f/<fid>")
def public_page(fid):
    return current_app.send_static_file("form.html")


@bp.get("/api/public/forms/<fid>")
def public_form(fid):
    f = forms().get(fid)
    if not f:
        raise ValueError("Ce questionnaire n'existe pas ou a été supprimé")
    return jsonify({"title": f["title"], "description": f["description"], "questions": f["questions"], "open": is_open(f),
                    "color": f["settings"].get("color", "#1f6bff")})


def validate(q, v):
    if v in (None, "", []):
        if q["required"]:
            raise ValueError(f"« {q['label']} » est obligatoire")
        return None
    t = q["type"]
    if t == "email" and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", str(v)):
        raise ValueError(f"« {q['label']} » : adresse e-mail invalide")
    if t == "number":
        try:
            return float(v) if "." in str(v) else int(v)
        except ValueError:
            raise ValueError(f"« {q['label']} » : nombre attendu") from None
    if t in ("scale", "rating"):
        v = int(v)
        lo, hi = (q.get("min", 1), q.get("max", 5)) if t == "scale" else (1, 5)
        if not lo <= v <= hi:
            raise ValueError(f"« {q['label']} » : valeur hors limites")
        return v
    if t == "yesno":
        return "Oui" if v in (True, "Oui", "oui", "1") else "Non"
    if t == "checkbox":
        return [str(x)[:500] for x in (v if isinstance(v, list) else [v])][:50]
    return str(v)[:5000]


@bp.post("/api/public/forms/<fid>")
def submit(fid):
    f = forms().get(fid)
    if not f:
        raise ValueError("Ce questionnaire n'existe pas")
    if not is_open(f):
        raise ValueError("Ce questionnaire est fermé")
    answers = (request.get_json(silent=True) or {}).get("answers") or {}
    clean = {q["id"]: validate(q, answers.get(q["id"])) for q in f["questions"]}
    store.append(f"forms/{fid}", {"id": secrets.token_hex(6), "at": time.time(), "answers": clean})
    if f["settings"].get("notify", True):
        activity.log(f["owner"], "questionnaires", f"Nouvelle réponse à « {f['title']} »", f"#/questionnaires/{fid}/reponses")
    return jsonify(ok=True, message=f["settings"].get("confirmation") or "Merci !")


def search(q, user):
    return [{"type": "form", "title": f["title"], "sub": "Questionnaire", "link": f"#/questionnaires/{f['id']}/modifier"}
            for f in forms().values() if f["owner"] == user and q in f["title"].lower()][:6]


def count(user):
    return sum(1 for f in forms().values() if f["owner"] == user)
